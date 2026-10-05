"""data/*.json → 예산서_게시현황.xlsx / .csv"""
import csv, json, glob, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
SEOUL = "종로구 중구 용산구 성동구 광진구 동대문구 중랑구 성북구 강북구 도봉구 노원구 은평구 서대문구 마포구 양천구 강서구 구로구 금천구 영등포구 동작구 관악구 서초구 강남구 송파구 강동구".split()
GG = "수원시 성남시 의정부시 안양시 부천시 광명시 평택시 동두천시 안산시 고양시 과천시 구리시 남양주시 오산시 시흥시 군포시 의왕시 하남시 용인시 파주시 이천시 안성시 김포시 화성시 광주시 양주시 포천시 여주시 연천군 가평군 양평군".split()
ORDER = [("서울특별시", g) for g in SEOUL] + [("경기도", g) for g in GG]
KIND = {"본예산": 0, "기금운용계획": 1, "추경": 2}

rows, summary = [], []
for sido, sgg in ORDER:
    p = os.path.join(BASE, "data", f"{sido}_{sgg}.json")
    if not os.path.exists(p):
        summary.append([sido, sgg, "", "", "자료 없음"] + [""] * 9)
        continue
    d = json.load(open(p, encoding="utf-8"))
    posts = sorted(d.get("게시물", []), key=lambda x: (x.get("연도") or 0, KIND.get(x.get("구분"), 9), x.get("회차") or 0))
    for x in posts:
        files = x.get("첨부파일") or []
        rows.append([sido, sgg, x.get("연도"), x.get("구분"), x.get("회차") or "", x.get("제목"),
                     x.get("등록일") or "", x.get("게시판URL") or "", x.get("게시물URL") or "",
                     "; ".join(f.get("파일명") or "" for f in files),
                     "; ".join(f.get("저장경로") or "" for f in files if f.get("저장경로")),
                     x.get("비고") or ""])
    cell = []
    for y in (2024, 2025, 2026):
        for k in ("본예산", "기금운용계획", "추경"):
            ps = [x for x in posts if x.get("연도") == y and x.get("구분") == k]
            if k == "추경":
                cell.append(", ".join(f"{x.get('회차') or '?'}회 {x.get('등록일') or '?'}" for x in ps))
            else:
                cell.append(", ".join(x.get("등록일") or "날짜미확인" for x in ps))
    boards = "\n".join(f"{b.get('이름','')}: {b.get('URL','')}" for b in d.get("게시판", []))
    summary.append([sido, sgg, d.get("홈페이지", ""), boards, d.get("비고", "")] + cell)

H_ROWS = ["시도", "지자체", "연도", "구분", "회차", "게시물 제목", "등록일", "게시판 URL", "게시물 URL", "첨부파일명", "다운로드 저장경로(2026)", "비고"]
H_SUM = ["시도", "지자체", "홈페이지", "예산 게시판", "비고"] + [f"{y} {k}" for y in (2024, 2025, 2026) for k in ("본예산", "기금운용계획", "추경")]

wb = Workbook()
for ws, head, data, widths in ((wb.active, H_SUM, summary, [10, 9, 28, 60, 40] + [14] * 9),
                               (wb.create_sheet(), H_ROWS, rows, [10, 9, 6, 11, 5, 40, 11, 50, 60, 40, 50, 40])):
    ws.append(head)
    for r in data:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
wb.worksheets[0].title = "요약(지자체별 등록일)"
wb.worksheets[1].title = "게시물 목록"
wb.save(os.path.join(BASE, "예산서_게시현황.xlsx"))
with open(os.path.join(BASE, "예산서_게시현황.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f); w.writerow(H_ROWS); w.writerows(rows)
print(f"지자체 {sum(1 for s in summary if s[4] != '자료 없음')}/{len(ORDER)}, 게시물 {len(rows)}건")
