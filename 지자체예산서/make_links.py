"""data/*.json → 다운로드링크_목록.xlsx / .csv (첨부파일 1개 = 1행)"""
import csv, json, os, re
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from build_index import ORDER, KIND, BASE

# 링크만으로는 안 받아지는 사이트 (수집 때 확인한 내용)
CAUTION = {
    "화성시": "POST 요청 + money_search_csrf 세션토큰 필요 — 게시판 페이지에서 직접 받을 것",
    "영등포구": "atchmnflStr 값이 요청마다 바뀌는 세션 토큰 — 링크가 만료되면 게시물URL에서 다시 받을 것",
    "강서구": "upperNo·fileNo가 페이지를 열 때마다 바뀜 — 링크가 안 되면 게시물URL에서 다시 받을 것",
    "강북구": "쿠키 기반 봇 확인(sabSignature) — 브라우저로 먼저 사이트를 연 뒤 받을 것",
}

HEAD = ["시도", "지자체", "연도", "구분", "회차", "게시물 제목", "등록일", "파일명", "다운로드URL",
        "요청방식·주의", "크기(바이트)", "2026 수집결과", "게시물URL", "게시판URL", "비고"]
rows = []
for sido, sgg in ORDER:
    p = os.path.join(BASE, "data", f"{sido}_{sgg}.json")
    d = json.load(open(p, encoding="utf-8"))
    posts = sorted(d["게시물"], key=lambda x: (x.get("연도") or 0, KIND.get(x.get("구분"), 9), x.get("회차") or 0))
    for x in posts:
        for f in x.get("첨부파일") or []:
            raw = f.get("다운로드URL") or ""
            m = re.match(r"(\S+)\s*\((.*)\)\s*$", raw)          # "URL (POST ...)" 형태 분리
            url, how = (m.group(1), m.group(2)) if m else (raw, "")
            how = "; ".join(s for s in (how, CAUTION.get(sgg, "")) if s)
            size = f.get("크기")
            if x.get("연도") == 2026:
                saved = f.get("저장경로")
                fail = re.search(r"빈 (HTML|응답)|0바이트", (f.get("비고") or "") + (x.get("비고") or ""))
                result = ("받음" if saved and os.path.exists(os.path.join(BASE, saved))
                          else "실패(서버 빈 응답)" if fail else "받지 않음(중복·대상 외)")
                if saved and size is None:
                    size = os.path.getsize(os.path.join(BASE, saved))
            else:
                result = ""
            note = "; ".join(s for s in (f.get("설명"), f.get("항목"), f.get("비고")) if s)
            rows.append([sido, sgg, x.get("연도"), x.get("구분"), x.get("회차") or "", x.get("제목"),
                         x.get("등록일") or "", f.get("파일명"), url, how, size, result,
                         x.get("게시물URL") or "", x.get("게시판URL") or "", note])

WIDTH = [10, 9, 6, 11, 5, 36, 11, 40, 70, 30, 12, 10, 50, 50, 30]
wb = Workbook()
for ws, data in ((wb.active, [r for r in rows if r[2] == 2026]), (wb.create_sheet(), rows)):
    ws.append(HEAD)
    for r in data:
        ws.append(r)
        c = ws.cell(ws.max_row, 9)
        if c.value and c.value.startswith("http"):
            c.hyperlink = c.value; c.style = "Hyperlink"
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
wb.worksheets[0].title = "2026년"
wb.worksheets[1].title = "전체(2024~2026)"
wb.save(os.path.join(BASE, "다운로드링크_목록.xlsx"))
with open(os.path.join(BASE, "다운로드링크_목록.csv"), "w", newline="", encoding="utf-8-sig") as fp:
    w = csv.writer(fp); w.writerow(HEAD); w.writerows(rows)
n26 = [r for r in rows if r[2] == 2026]
print(f"전체 {len(rows)}개, 2026 {len(n26)}개 {dict(Counter(r[11] for r in n26))}, 주의 {sum(bool(r[9]) for r in rows)}")
